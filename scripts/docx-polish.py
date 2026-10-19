"""Post-process the LibreOffice output so the .docx behaves like a printed
handover document rather than a web page export: A4 page setup, print margins,
a Chinese-capable body font, and repeated table header rows across pages.
"""
import pathlib
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = pathlib.Path('/Users/sumuzhi/WorkBuddy/2026-10-19-14-44-12/build/delivery.docx')
DST = pathlib.Path('/Users/sumuzhi/WorkBuddy/2026-10-19-14-44-12/Ostra-交付文档-v1.0.docx')

doc = Document(SRC)

# --- page setup -----------------------------------------------------------
for section in doc.sections:
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(18)
    section.left_margin = Mm(16)
    section.right_margin = Mm(16)


def set_fonts(run_or_style, latin='Inter', east='PingFang SC', size=None, bold=None):
    font = run_or_style.font
    font.name = latin
    if size:
        font.size = size
    if bold is not None:
        font.bold = bold
    rpr = run_or_style.element.get_or_add_rPr() if hasattr(run_or_style.element, 'get_or_add_rPr') else None
    if rpr is not None:
        rfonts = rpr.find(qn('w:rFonts'))
        if rfonts is None:
            rfonts = OxmlElement('w:rFonts')
            rpr.append(rfonts)
        rfonts.set(qn('w:eastAsia'), east)
        rfonts.set(qn('w:ascii'), latin)
        rfonts.set(qn('w:hAnsi'), latin)


# --- body + heading fonts -------------------------------------------------
styles = doc.styles
for name, size, bold, color in [
    ('Normal', 10.5, False, None),
    ('Heading 1', 22, True, RGBColor(0x1B, 0x19, 0x16)),
    ('Heading 2', 14, True, RGBColor(0x1B, 0x19, 0x16)),
    ('Heading 3', 11.5, True, RGBColor(0x57, 0x50, 0x4A)),
]:
    try:
        st = styles[name]
    except KeyError:
        continue
    set_fonts(st, size=Pt(size), bold=bold)
    if color is not None:
        st.font.color.rgb = color
    st.paragraph_format.space_before = Pt(0 if name == 'Normal' else 12)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.5 if name == 'Normal' else 1.2

# --- tables: repeat header row, keep rows unbreakable ----------------------
for table in doc.tables:
    if not table.rows:
        continue
    hdr = table.rows[0]
    tr_pr = hdr._tr.get_or_add_trPr()
    if tr_pr.find(qn('w:tblHeader')) is None:
        el = OxmlElement('w:tblHeader')
        el.set(qn('w:val'), 'true')
        tr_pr.append(el)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn('w:cantSplit')) is None:
            tr_pr.append(OxmlElement('w:cantSplit'))

# --- title: give the cover line room --------------------------------------
if doc.paragraphs and doc.paragraphs[0].text.strip():
    p = doc.paragraphs[0]
    p.paragraph_format.space_after = Pt(12)
    for run in p.runs:
        run.bold = True

doc.save(DST)
print('saved:', DST, DST.stat().st_size // 1024, 'KB')
print('paragraphs:', len(doc.paragraphs), 'tables:', len(doc.tables))
